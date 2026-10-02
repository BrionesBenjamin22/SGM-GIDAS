import os
from pathlib import Path
import sqlite3
import subprocess
import sys
import tempfile
import unittest
from datetime import date
from decimal import Decimal
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook

from extension import db
from modules import models_registry  # noqa: F401
from modules.auth.models.usuario import RolUsuario, Usuario
from modules.grupo.models.grupo import GrupoInvestigacionUtn
from modules.informes.models.informe import Informe
from modules.memorias.models.memorias import Memoria
from modules.memorias.services.memoria_service import MemoriaService
from modules.memorias.services.exportacion_service_impl import ExportService
from modules.personal.models.personal import Investigador
from modules.recursos.models.movimiento_financiero import MovimientoFinanciero
from tools.memoria_2025_dataset import amount, first_date, read_dataset
from tools.seed_memoria_2025 import seed_memoria_2025


class Memoria2025SeedTest(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(APP_ENV="testing", TESTING=True, SQLALCHEMY_DATABASE_URI="sqlite://",
                               SQLALCHEMY_TRACK_MODIFICATIONS=False)
        db.init_app(self.app)
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        rol = RolUsuario(nombre="ADMIN")
        db.session.add(rol)
        db.session.flush()
        self.admin = Usuario(nombre_usuario="admin", mail="admin@gidas.local", contrasena="unused", id_rol=rol.id)
        db.session.add(self.admin)
        db.session.commit()
        self.addCleanup(self.cleanup)

    def cleanup(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def test_extraccion_fechas_importes_y_contenido_real(self):
        data = read_dataset()
        self.assertEqual(len(data["investigadores"]), 6)
        self.assertEqual(len(data["becarios"]), 23)
        self.assertEqual(len(data["equipamiento"]), 3)
        self.assertEqual(len(data["reuniones"]), 6)
        self.assertEqual(sum(amount(row["ingresos"]) for row in data["finanzas"]), Decimal("9195000.00"))
        self.assertEqual(sum(amount(row["egresos"]) for row in data["finanzas"]), Decimal("9195000.00"))
        self.assertEqual(first_date(data["equipamiento"][0]["fecha"]), date(2025, 10, 17))
        self.assertEqual(first_date("13 y 14 agosto 2025"), date(2025, 8, 13))
        self.assertEqual(first_date("diciembre 2025"), date(2025, 12, 1))
        self.assertEqual(data["formulas"]["C309"], "=1295000+700000")

    def test_seed_idempotente_cierre_exportacion_e_historia(self):
        result = seed_memoria_2025(self.admin.id)
        counts = {table.name: db.session.query(table).count() for table in db.metadata.tables.values()}
        repeated = seed_memoria_2025(self.admin.id)
        self.assertFalse(repeated["created"])
        self.assertEqual(counts, {table.name: db.session.query(table).count() for table in db.metadata.tables.values()})
        self.assertEqual(MovimientoFinanciero.query.count(), 12)
        memoria = db.session.get(Memoria, result["memoria_id"])
        MemoriaService.change_status(memoria.id, {"estado": "cerrada"}, self.admin.id)
        exported = ExportService.generar_excel_memoria(memoria.id, memoria.version_actual_id)
        workbook = load_workbook(exported)
        self.assertEqual(workbook.sheetnames, ["Hoja1", "Hoja2"])
        cells = list(workbook["Hoja1"].values)
        texts = [str(value) for row in cells for value in row if value is not None]
        self.assertIn("MEMORIAS 2025", texts[0])
        finance_start = next(i for i, row in enumerate(cells) if row[0] == "Erogaciones Corrientes")
        finance_end = next(i for i, row in enumerate(cells) if "VI - PROGRAMA" in str(row[0]))
        finance = cells[finance_start:finance_end]
        self.assertEqual(sum(row[2] for row in finance if isinstance(row[2], (float, int))), 9195000)
        self.assertEqual(sum(row[4] for row in finance if isinstance(row[4], (float, int))), 9195000)
        self.assertFalse(any(cell.data_type == "f" for row in workbook["Hoja1"] for cell in row))
        self.assertTrue(memoria.version_actual.contexto_institucional["logros_proyectos"])
        self.assertTrue(any("Leandro ROCCA" in value for value in texts))
        self.assertGreater(len(MemoriaService.get_becarios_snapshot(memoria.id, memoria.version_actual_id)), 0)
        if os.getenv("MEMORIA_VALIDATION_OUTPUT"):
            workbook.save(Path(os.environ["MEMORIA_VALIDATION_OUTPUT"]))
        for informe in Informe.query.all():
            informe.resultados = "Resultado actual cambiado"
            informe.actividades = "Actividad actual cambiada"
        grupo = db.session.get(GrupoInvestigacionUtn, result["grupo_id"])
        grupo.objetivo_desarrollo = "Objetivo actual cambiado"
        Investigador.query.first().nombre_apellido = "Nombre actual cambiado"
        db.session.commit()
        after = load_workbook(ExportService.generar_excel_memoria(memoria.id, memoria.version_actual_id))
        self.assertEqual(cells, list(after["Hoja1"].values))

    def test_rechaza_produccion_y_uct_ajena_sin_cambios(self):
        self.app.config["APP_ENV"] = "production"
        with self.assertRaises(RuntimeError):
            seed_memoria_2025(self.admin.id)
        self.app.config["APP_ENV"] = "testing"
        db.session.add(GrupoInvestigacionUtn(nombre_sigla_grupo="Otro grupo", nombre_unidad_academica="Otra FR", mail="grupo@utn.local", objetivo_desarrollo="Objetivo"))
        db.session.commit()
        with self.assertRaises(RuntimeError):
            seed_memoria_2025(self.admin.id)
        self.assertEqual(GrupoInvestigacionUtn.query.count(), 1)

    def test_exportador_no_recorta_y_no_ejecuta_formulas_de_texto(self):
        result = seed_memoria_2025(self.admin.id)
        memoria = db.session.get(Memoria, result["memoria_id"])
        MemoriaService.change_status(memoria.id, {"estado": "cerrada"}, self.admin.id)
        sources = ExportService._build_memoria_snapshot_sources(memoria.id, memoria.version_actual_id)
        sources["investigadores"] = [{"nombre_apellido": f"Investigador {index}", "horas_semanales": 10} for index in range(25)]
        sources["becarios"] = [{"nombre_apellido": "Doctorando", "tipo_formacion_nombre": "Doctorado", "horas_semanales": 20}]
        sources["articulos"] = [{"titulo": "Contenido", "descripcion": "=1+1"}]
        with patch.object(ExportService, "_build_memoria_snapshot_sources", return_value=sources):
            workbook = load_workbook(ExportService.generar_excel_memoria(memoria.id, memoria.version_actual_id))
        cells = [cell for row in workbook["Hoja1"] for cell in row]
        values = [cell.value for cell in cells]
        for index in range(25):
            self.assertIn(f"Investigador {index}", values)
        self.assertIn("Doctorando", values)
        self.assertEqual(next(cell for cell in cells if cell.value == "=1+1").data_type, "s")
        self.assertTrue(any("VI - PROGRAMA" in str(value) for value in values))

    def test_cli_inicializa_usuarios_roles_y_excel_en_base_dedicada(self):
        with tempfile.TemporaryDirectory() as directory:
            database = Path(directory) / "memoria.sqlite"
            output = Path(directory) / "memoria.xlsx"
            environment = dict(os.environ, APP_ENV="testing", ENV_FILE=str(Path(directory) / "absent.env"),
                DATABASE_URL="sqlite:///" + database.as_posix(), SEED_ADMIN_PASSWORD="Memorias2025!", LOG_LEVEL="CRITICAL")
            script = Path(__file__).resolve().parents[1] / "tools" / "seed_memoria_2025.py"
            completed = subprocess.run([sys.executable, str(script), "--apply", "--initialize", "--close", "--output", str(output)],
                env=environment, capture_output=True, text=True, timeout=90)
            self.assertEqual(completed.returncode, 0, completed.stderr)
            self.assertTrue(output.is_file())
            with sqlite3.connect(database) as connection:
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM usuario_grupo_utn").fetchone()[0], 3)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM memoria_version WHERE estado='cerrada'").fetchone()[0], 1)
                self.assertEqual(connection.execute("SELECT COUNT(*) FROM movimiento_financiero").fetchone()[0], 12)


if __name__ == "__main__":
    unittest.main()
