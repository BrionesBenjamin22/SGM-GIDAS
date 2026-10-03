from copy import copy
from datetime import date
from types import SimpleNamespace
import unittest

from openpyxl import load_workbook

from modules.memorias.services.exportacion_service_impl import ExportService
from modules.memorias.services.memoria_excel_template import render_memoria, CONTENT_FONT


class MemoriaFormatoTest(unittest.TestCase):
    def render(self, **records):
        sources = {key: [] for key in (
            "investigadores", "personal", "becarios", "equipamiento", "documentacion",
            "proyectos", "distinciones", "participaciones", "visitas", "trabajos_reunion",
            "trabajos_revista", "articulos", "registros", "actividades", "transferencias", "erogaciones",
        )}
        sources.update(records, memoria=SimpleNamespace(periodo_inicio=date(2026, 1, 1), periodo_fin=date(2026, 12, 31)))
        context = {"grupo": {"nombre_sigla_grupo": "GIDAS", "nombre_unidad_academica": "UTN",
                             "mail": "gidas@utn.edu.ar", "objetivo_desarrollo": "Objetivo institucional"}}
        original = load_workbook(ExportService.MEMORIA_TEMPLATE_PATH)
        generated = load_workbook(render_memoria(load_workbook(ExportService.MEMORIA_TEMPLATE_PATH), sources, context, ExportService))
        return original, generated

    def test_maquetacion_completa_sin_registros_respeta_referencia(self):
        original, generated = self.render()
        self.assertEqual(original.sheetnames, generated.sheetnames)
        for name in original.sheetnames:
            source, output = original[name], generated[name]
            self.assertEqual(source.calculate_dimension(), output.calculate_dimension())
            self.assertEqual(set(map(str, source.merged_cells.ranges)), set(map(str, output.merged_cells.ranges)))
            self.assertEqual(dict(source.page_setup), dict(output.page_setup))
            self.assertEqual(source.page_margins, output.page_margins)
            self.assertEqual(source.print_options, output.print_options)
            self.assertEqual(source.print_area, output.print_area)
            self.assertEqual(source.freeze_panes, output.freeze_panes)
            self.assertEqual(dict(source.sheet_format), dict(output.sheet_format))
            self.assertEqual({key: dict(value) for key, value in source.row_dimensions.items()},
                             {key: dict(value) for key, value in output.row_dimensions.items()})
            self.assertEqual({key: dict(value) for key, value in source.column_dimensions.items()},
                             {key: dict(value) for key, value in output.column_dimensions.items()})
            for row in source:
                for cell in row:
                    target = output[cell.coordinate]
                    for property_name in ("font", "fill", "border", "alignment", "number_format", "protection"):
                        if property_name == "font" and target.value is not None and target.font == CONTENT_FONT:
                            continue  # El contenido se normaliza; encabezados y estilos restantes se conservan.
                        self.assertEqual(copy(getattr(cell, property_name)), copy(getattr(target, property_name)),
                                         f"{cell.coordinate}: {property_name}")
        source, output = original["Hoja1"], generated["Hoja1"]
        for row in (3, 4, 10, 17, 19, 21, 22, 23, 25, 27, 28, 29, 30, 32, 34, 35, 36, 37,
                    39, 41, 42, 44, 45, 59, 61, 70, 71, 75, 76, 78, 79, 80, 84, 85, 91, 92,
                    96, 97, 106, 107, 118, 119, 124, 125, 130, 131, 135, 136, 137, 138, 143,
                    146, 148, 149, 191, 196, 205, 209, 210, 211, 219, 220, 224, 225, 226, 232,
                    238, 250, 256, 257, 261, 265, 266, 274, 275, 276, 277, 282, 283, 287,
                    288, 291, 292, 296, 297, 300, 301, 305, 306, 307, 308, 316, 317):
            self.assertEqual(list(source.values)[row - 1], list(output.values)[row - 1], f"Encabezado {row}")
        self.assertIsNone(output["A2"].value)
        self.assertIsNone(output["H221"].value)
        self.assertFalse(any(cell.data_type == "f" for row in output for cell in row))

    def test_desborde_conserva_encabezados_posteriores_y_estilo_de_filas(self):
        original, generated = self.render(investigadores=[{"nombre_apellido": f"Persona {i}"} for i in range(25)])
        source, output = original["Hoja1"], generated["Hoja1"]
        extra = 25 - 6
        self.assertEqual(output.max_row, source.max_row + extra)
        self.assertEqual(output["B86"].value, "Persona 24")
        for original_row in (70, 124, 135, 209, 274, 305, 307, 316):
            new_row = original_row + extra
            self.assertEqual(source[f"A{original_row}"].value, output[f"A{new_row}"].value)
            self.assertEqual(source.row_dimensions[original_row].height, output.row_dimensions[new_row].height)
            self.assertEqual(copy(source[f"A{original_row}"].alignment), copy(output[f"A{new_row}"].alignment))
        for new_row in range(68, 87):
            self.assertEqual(CONTENT_FONT, copy(output[f"B{new_row}"].font))
            self.assertEqual(copy(source["B62"].border), copy(output[f"B{new_row}"].border))

    def test_varios_desbordes_y_doctorandos_preservan_el_resto_del_formato(self):
        original, generated = self.render(
            personal=[{"nombre_apellido": f"Apoyo {i}", "tipo_personal_nombre": "Apoyo"} for i in range(6)],
            becarios=[{"nombre_apellido": f"Doctorando {i}", "tipo_formacion_nombre": "Doctorado"} for i in range(5)],
            equipamiento=[{"denominacion": f"Equipo {i}"} for i in range(16)],
            proyectos=[{"codigo_proyecto": f"PID-{i}"} for i in range(6)],
        )
        source, output = original["Hoja1"], generated["Hoja1"]
        offsets = ((78, 5), (84, 2), (129, 13), (142, 3))
        for row in source:
            for cell in row:
                if 81 <= cell.row <= 83:  # Reemplaza el bloque combinado "No aplica" por filas de doctorandos
                    continue
                target_row = cell.row + sum(amount for start, amount in offsets if cell.row >= start)
                target = output.cell(target_row, cell.column)
                for property_name in ("font", "fill", "border", "alignment", "number_format"):
                    if property_name == "font" and target.value is not None and target.font == CONTENT_FONT:
                        continue
                    self.assertEqual(copy(getattr(cell, property_name)), copy(getattr(target, property_name)),
                                     f"{cell.coordinate} -> {target.coordinate}: {property_name}")
        for index in range(5):
            row = 86 + index
            self.assertEqual(output[f"B{row}"].value, f"Doctorando {index}")
            self.assertEqual(CONTENT_FONT, copy(output[f"B{row}"].font))
            self.assertEqual(copy(source["B86"].border), copy(output[f"B{row}"].border))
        self.assertEqual(output["B148"].value, "Equipo 15")
        self.assertEqual(output["C164"].value, "PID-5")
        self.assertEqual(source["A307"].value, output["A330"].value)

    def test_textos_extensos_fuente_uniforme_y_visitantes_en_filas_propias(self):
        original, generated = self.render(
            investigadores=[{"nombre_apellido": "Camila Soledad Benítez", "categoria_utn_nombre": "Investigador formado"}],
            articulos=[{"descripcion": "Descripción académica de resultados. " * 60}],
            visitas=[{"procedencia": f"Institución {index}", "razon": "Cooperación científica. " * 12} for index in range(5)],
        )
        source, output = original["Hoja1"], generated["Hoja1"]
        self.assertEqual(output["B62"].font.sz, 11)
        self.assertEqual(output["B62"].font.color.rgb, "FF000000")
        for index in range(5):
            self.assertTrue(output[f"A{149 + index}"].value.startswith(f"Institución {index}:"))
            self.assertEqual(CONTENT_FONT, copy(output[f"A{149 + index}"].font))
        self.assertEqual(source["A209"].value, output["A214"].value)
        self.assertGreater(output.row_dimensions[245].height, source.row_dimensions[240].height or 15)
        self.assertEqual(copy(source["A209"].font), copy(output["A214"].font))

    def test_varios_parrafos_ocupan_filas_separadas_sin_compactar_secciones(self):
        _, generated = self.render(participaciones=[
            {"nombre_evento": f"Curso de actualización {index}", "forma_participacion": "Contenido académico. " * 25}
            for index in range(6)
        ])
        output = generated["Hoja1"]
        for index in range(6):
            self.assertTrue(output[f"A{192 + index}"].value.startswith(f"Curso de actualización {index}"))
            self.assertEqual(CONTENT_FONT, copy(output[f"A{192 + index}"].font))
            self.assertLess(output.row_dimensions[192 + index].height or 15, 409)
        self.assertIn("GESTIÓN", str(output["A198"].value).upper())


if __name__ == "__main__":
    unittest.main()
