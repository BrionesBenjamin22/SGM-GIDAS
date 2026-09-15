import unittest
import importlib.util
from pathlib import Path
from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch

from flask import Flask
from openpyxl import load_workbook
import sqlalchemy as sa
from alembic.migration import MigrationContext
from alembic.operations import Operations

from extension import db
from modules import models_registry  # noqa: F401
from modules.grupo.models.grupo import GrupoInvestigacionUtn
from modules.grupo.models.programa_actividades import PlanificacionGrupo
from modules.memorias.models.memorias import Memoria
from modules.memorias.routes.memorias_rutas import memoria_bp
from modules.memorias.services.memoria_service import MemoriaService
from modules.memorias.services.exportacion_service_impl import ExportService
from modules.personal.models.personal import Investigador, InvestigadorHorasHistorial
from modules.produccion.models.articulo_divulgacion import ArticuloDivulgacion
from modules.produccion.models.distinciones import DistincionRecibida
from modules.proyectos.models.participacion_relevante import ParticipacionRelevante
from modules.proyectos.models.proyecto_investigacion import ProyectoInvestigacion, TipoProyecto
from modules.shared.exceptions import ConflictError, NotFoundError, ValidationError
from modules.shared.models.auditoria_campo import AuditoriaCampo


class MemoriaPeriodosUctTest(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SQLALCHEMY_DATABASE_URI="sqlite:///:memory:")
        db.init_app(self.app)
        self.app.register_blueprint(memoria_bp, url_prefix="/api/v1/memorias")
        self.context = self.app.app_context()
        self.context.push()
        db.create_all()
        db.session.add_all([
            GrupoInvestigacionUtn(id=n, nombre_sigla_grupo=f"UCT {n}", mail=f"uct{n}@test.invalid",
                                  nombre_unidad_academica="Regional", objetivo_desarrollo="Investigación")
            for n in (1, 2)
        ])
        db.session.commit()
        self.client = self.app.test_client()
        self.headers = {"Authorization": "Bearer test"}

    def tearDown(self):
        db.session.remove()
        db.drop_all()
        self.context.pop()

    def crear(self, grupo=1, inicio="2025-07-01", fin="2026-06-30"):
        return MemoriaService.create({"grupo_utn_id": grupo, "periodo_inicio": inicio, "periodo_fin": fin}, 1)

    def cerrar(self, memoria):
        return MemoriaService.change_status(memoria["id"], {"estado": "cerrada"}, 1)

    def test_matriz_permisos_alta_y_correccion(self):
        for rol, status in (("LECTURA", 403), ("LECTOR", 403), ("GESTOR", 201), ("ADMIN", 201)):
            with self.subTest(rol=rol), patch("modules.shared.services.middleware.AuthService.verify_token",
                                            return_value={"sub": "1", "rol": rol}):
                response = self.client.post("/api/v1/memorias", headers=self.headers, json={
                    "grupo_utn_id": 1 if rol == "GESTOR" else 2,
                    "periodo_inicio": "2025-07-01", "periodo_fin": "2026-06-30",
                })
                self.assertEqual(response.status_code, status)
                if status == 201:
                    memoria = response.get_json()
                    edited = self.client.put(f"/api/v1/memorias/{memoria['id']}", headers=self.headers,
                                             json={"periodo_fin": "2026-07-31"})
                    self.assertEqual(edited.status_code, 200)
                    history = self.client.get(f"/api/v1/memorias/{memoria['id']}/historial", headers=self.headers)
                    self.assertEqual(history.get_json()[0]["valor_nuevo"], "2026-07-31")
                    changed_state = self.client.put(
                        f"/api/v1/memorias/{memoria['id']}/estado",
                        headers=self.headers,
                        json={"estado": "en revision"},
                    )
                    self.assertEqual(changed_state.status_code, 200)
                else:
                    self.assertEqual(self.client.put("/api/v1/memorias/1", headers=self.headers,
                                                     json={"periodo_fin": "2026-07-31"}).status_code, 403)
                    self.assertEqual(self.client.put("/api/v1/memorias/1/estado", headers=self.headers,
                                                     json={"estado": "en revision"}).status_code, 403)
        self.assertEqual(self.client.post("/api/v1/memorias", json={}).status_code, 401)

    def test_periodos_validos_y_uct_requerida(self):
        for inicio, fin in (("2009-12-31", "2026-01-01"), ("2026-08-01", "2026-07-01"),
                            ("invalida", "2026-01-01")):
            with self.subTest(inicio=inicio), self.assertRaises(ValidationError):
                self.crear(inicio=inicio, fin=fin)
        with self.assertRaises(ValidationError):
            self.crear(grupo=999)
        with self.assertRaises(ValidationError):
            self.crear(grupo=None)
        self.assertEqual(self.crear()["periodo_fin"], "2026-06-30")

    def test_solapamiento_inclusivo_y_periodos_consecutivos_mismo_anio(self):
        primera = self.crear(inicio="2026-01-01", fin="2026-03-31")
        self.cerrar(primera)
        for inicio, fin in (("2026-03-31", "2026-04-30"), ("2026-02-01", "2026-02-28"),
                            ("2025-12-01", "2026-12-31")):
            with self.subTest(inicio=inicio), self.assertRaises(ConflictError):
                self.crear(inicio=inicio, fin=fin)
        segunda = self.crear(inicio="2026-04-01", fin="2026-06-30")
        self.assertNotEqual(primera["id"], segunda["id"])

    def test_activa_por_uct_y_baja_no_bloquea(self):
        primera = self.crear()
        otra = self.crear(grupo=2)
        self.assertEqual(otra["grupo_utn_id"], 2)
        with self.assertRaises(ConflictError):
            self.crear(inicio="2027-01-01", fin="2027-12-31")
        MemoriaService.delete(primera["id"], 1)
        self.assertEqual(self.crear()["grupo_utn_id"], 1)

    def test_correccion_solo_diferencias_historial_y_bloqueo_tras_reapertura(self):
        memoria = self.crear()
        updated = MemoriaService.update(memoria["id"], {"periodo_inicio": "2025-08-01"}, 2)
        self.assertEqual(updated["updated_by"], 2)
        history = MemoriaService.get_historial(memoria["id"])
        self.assertEqual(len(history), 1)
        self.assertEqual(history[0]["valor_anterior"], "2025-07-01")
        self.assertEqual(history[0]["campo"], "periodo_inicio")
        with patch.object(db.session, "commit") as commit:
            MemoriaService.update(memoria["id"], {"periodo_inicio": "2025-08-01"}, 2)
            commit.assert_not_called()
        self.cerrar(memoria)
        MemoriaService.reopen(memoria["id"], 1)
        with self.assertRaises(ConflictError):
            MemoriaService.update(memoria["id"], {"periodo_fin": "2027-06-30"}, 2)
        history = MemoriaService.get_historial(memoria["id"])
        self.assertEqual(len(history), 3)
        self.assertEqual({item["campo"] for item in history}, {"periodo_inicio", "estado", "versiones"})

    def test_uct_inmutable_y_asociacion_explicita_memoria_anterior(self):
        memoria = self.crear(grupo=2)
        with self.assertRaises(ConflictError):
            MemoriaService.update(memoria["id"], {"grupo_utn_id": 1}, 1)
        anterior = Memoria(periodo_inicio=date(2024, 1, 1), periodo_fin=date(2024, 12, 31), created_by=1)
        db.session.add(anterior)
        db.session.flush()
        MemoriaService._crear_version_inicial(anterior, 1)
        db.session.commit()
        with self.assertRaises(ConflictError):
            MemoriaService.change_status(anterior.id, {"estado": "cerrada"}, 1)
        self.assertEqual(MemoriaService.update(anterior.id, {"grupo_utn_id": 1}, 1)["grupo_utn_id"], 1)
        historial = MemoriaService.get_historial(anterior.id)
        self.assertIsNone(historial[0]["valor_anterior"])
        self.assertEqual(historial[0]["valor_nuevo"], "UCT 1")
        evento = AuditoriaCampo.query.filter_by(
            entidad="memoria",
            registro_id=anterior.id,
            campo="grupo_utn_id",
        ).one()
        evento.valor_nuevo = 1
        db.session.commit()
        self.assertEqual(
            MemoriaService.get_historial(anterior.id)[0]["valor_nuevo"],
            "UCT 1",
        )

    def test_snapshot_filtra_uct_fechas_intervalos_horas_y_preserva_historia(self):
        db.session.add(TipoProyecto(id=1, nombre="PID"))
        plan = PlanificacionGrupo(grupo_id=1, anio=2027, descripcion="Plan original")
        db.session.add(plan)
        for n in (1, 2):
            investigador = Investigador(id=n, nombre_apellido=f"Investigador {n}", horas_semanales=40,
                                         fecha_alta_grupo=date(2020, 1, 1), grupo_utn_id=n)
            db.session.add(investigador)
            db.session.add_all([
                InvestigadorHorasHistorial(investigador=investigador, horas_semanales=15,
                                           fecha_inicio=date(2020, 1, 1), fecha_fin=date(2026, 6, 30)),
                InvestigadorHorasHistorial(investigador=investigador, horas_semanales=40,
                                           fecha_inicio=date(2026, 7, 1)),
                ParticipacionRelevante(nombre_evento=f"Evento {n}", forma_participacion="Ponente",
                                       fecha=date(2026, 1, 1), investigador=investigador),
            ])
            for fecha in (date(2025, 6, 30), date(2025, 7, 1), date(2025, 12, 31),
                          date(2026, 1, 1), date(2026, 6, 30), date(2026, 7, 1)):
                db.session.add(ArticuloDivulgacion(titulo=f"Artículo {n} {fecha}", descripcion="Texto",
                                                  fecha_publicacion=fecha, grupo_utn_id=n))
            proyecto = ProyectoInvestigacion(nombre_proyecto=f"Proyecto {n}", codigo_proyecto=f"PID{n}",
                                              descripcion_proyecto="Texto", tipo_proyecto_id=1, grupo_utn_id=n,
                                              fecha_inicio=date(2024, 1, 1), fecha_fin=date(2026, 1, 1))
            db.session.add(proyecto)
            db.session.add(DistincionRecibida(descripcion=f"Distinción {n}", fecha=date(2026, 1, 1),
                                             proyecto_investigacion=proyecto))
        db.session.add(ProyectoInvestigacion(nombre_proyecto="Terminado", codigo_proyecto="OLD", descripcion_proyecto="Texto",
                                              tipo_proyecto_id=1, grupo_utn_id=1, fecha_inicio=date(2024, 1, 1),
                                              fecha_fin=date(2025, 6, 30)))
        db.session.commit()
        primera = self.crear()
        segunda = self.crear(grupo=2)
        self.cerrar(primera)
        self.cerrar(segunda)
        version = primera["version_actual"]["id"]
        articulos = MemoriaService.get_articulos_divulgacion_snapshot(primera["id"], version)
        self.assertEqual(len(articulos), 4)
        self.assertEqual({item["grupo_utn_id"] for item in articulos}, {1})
        proyectos = MemoriaService.get_proyectos_snapshot(primera["id"], version)
        self.assertEqual([item["codigo_proyecto"] for item in proyectos], ["PID1"])
        integrantes = MemoriaService.get_investigadores_snapshot(primera["id"], version)
        self.assertEqual(integrantes[0]["horas_semanales"], 15)
        self.assertEqual(integrantes[0]["fecha_alta_grupo"], "2020-01-01")
        self.assertEqual(len(MemoriaService.get_participaciones_relevantes_snapshot(primera["id"], version)), 1)
        self.assertEqual(len(MemoriaService.get_distinciones_snapshot(primera["id"], version)), 1)
        with self.assertRaises(NotFoundError):
            MemoriaService.get_articulos_divulgacion_snapshot(segunda["id"], version)
        original = articulos[0]["titulo"]
        db.session.get(ArticuloDivulgacion, articulos[0]["articulo_divulgacion_id"]).titulo = "Título actual cambiado"
        db.session.commit()
        self.assertEqual(MemoriaService.get_articulos_divulgacion_snapshot(primera["id"], version)[0]["titulo"], original)
        content = ExportService.generar_excel_memoria(primera["id"], version)
        workbook = load_workbook(BytesIO(content) if isinstance(content, bytes) else content)
        self.assertIn("01/07/2025", workbook["Hoja1"]["A1"].value)
        self.assertIn("30/06/2026", workbook["Hoja1"]["A1"].value)
        self.assertIn("UCT 1", workbook["Hoja1"]["A1"].value)
        original_cells = list(workbook["Hoja1"].values)
        db.session.get(GrupoInvestigacionUtn, 1).nombre_sigla_grupo = "Nombre actual modificado"
        plan.descripcion = "Plan actual modificado"
        db.session.get(Investigador, 1).horas_semanales = 99
        db.session.commit()
        after = load_workbook(ExportService.generar_excel_memoria(primera["id"], version))
        self.assertEqual(list(after["Hoja1"].values), original_cells)

    def test_falla_snapshot_revierte_estado_y_fotos_parciales(self):
        memoria = self.crear()
        db.session.add(ArticuloDivulgacion(titulo="Original", descripcion="Texto", fecha_publicacion=date(2026, 1, 1), grupo_utn_id=1))
        db.session.commit()
        with patch("modules.memorias.services.memoria_service.snapshot_visitas_para_memoria_version", side_effect=RuntimeError("Fallo")):
            with self.assertRaises(RuntimeError):
                self.cerrar(memoria)
        root = db.session.get(Memoria, memoria["id"])
        self.assertEqual(root.version_actual.estado.value, "abierta")
        self.assertIsNone(root.version_actual.fecha_cierre)
        self.assertEqual(MemoriaService.get_articulos_divulgacion_snapshot(root.id, root.version_actual.id), [])


class MigracionMemoriaPeriodosTest(unittest.TestCase):
    def test_upgrade_downgrade_conserva_datos_sin_inventar_uct(self):
        path = Path(__file__).resolve().parents[1] / "migrations/versions/e16a0b2c4d60_memoria_periodos_uct.py"
        spec = importlib.util.spec_from_file_location("iss16_migration", path)
        migration = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(migration)
        engine = sa.create_engine("sqlite:///:memory:")
        self.addCleanup(engine.dispose)
        with engine.begin() as conn:
            conn.exec_driver_sql("CREATE TABLE grupo_utn (id INTEGER PRIMARY KEY)")
            conn.exec_driver_sql("CREATE TABLE memoria (id INTEGER PRIMARY KEY)")
            conn.exec_driver_sql("INSERT INTO memoria VALUES (1)")
            conn.exec_driver_sql("CREATE TABLE memoria_version (id INTEGER PRIMARY KEY)")
            conn.exec_driver_sql("INSERT INTO memoria_version VALUES (1)")
            for tabla in migration.TABLAS:
                conn.exec_driver_sql(f"CREATE TABLE {tabla} (id INTEGER PRIMARY KEY, horas_semanales INTEGER NOT NULL)")
                conn.exec_driver_sql(f"INSERT INTO {tabla} VALUES (1, 20)")
            with patch.object(migration, "op", Operations(MigrationContext.configure(conn))):
                migration.upgrade()
                self.assertIsNone(conn.exec_driver_sql("SELECT grupo_utn_id FROM memoria").scalar())
                self.assertIsNone(conn.exec_driver_sql("SELECT contexto_institucional FROM memoria_version").scalar())
                for tabla in migration.TABLAS:
                    self.assertIsNone(conn.exec_driver_sql(f"SELECT fecha_alta_grupo FROM {tabla}").scalar())
                    self.assertEqual(conn.exec_driver_sql(f"SELECT horas_semanales FROM {tabla}").scalar(), 20)
                migration.downgrade()
                self.assertNotIn("grupo_utn_id", [c["name"] for c in sa.inspect(conn).get_columns("memoria")])
                self.assertEqual(conn.exec_driver_sql("SELECT id FROM memoria").scalar(), 1)


if __name__ == "__main__":
    unittest.main()
