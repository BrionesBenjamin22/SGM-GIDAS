"""Lectura del documento institucional, sin acceso a la base de datos."""
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
import re

from openpyxl import load_workbook


SOURCE = Path(__file__).resolve().parents[1] / "assets" / "Memorias 2025 - GIDAS.xlsx"


def text(value):
    return str(value).strip() if value is not None else ""


def amount(value):
    if value in (None, "", "-"):
        return Decimal("0.00")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    raw = re.sub(r"[^\d,.]", "", str(value))
    if "," in raw and "." in raw:
        raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
    elif "," in raw:
        parts = raw.split(",")
        raw = "".join(parts) if len(parts[-1]) == 3 else raw.replace(",", ".")
    return Decimal(raw).quantize(Decimal("0.01"))


def first_date(value, fallback=date(2025, 1, 1)):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = text(value)
    match = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if match:
        return date(int(match[3]), int(match[2]), int(match[1]))
    months = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
              "julio": 7, "agosto": 8, "septiembre": 9, "octubre": 10, "noviembre": 11, "diciembre": 12}
    for name, month in months.items():
        if name in raw.lower():
            days = re.search(r"\b(\d{1,2})(?:\s*(?:,|y)\s*\d{1,2})*\s+(?:de\s+)?" + name + r"\b", raw.lower())
            year = re.search(r"20\d{2}", raw)
            return date(int(year[0]) if year else 2025, month, int(days[1]) if days else 1)
    return fallback


def read_dataset(path=SOURCE):
    """Mantiene coordenadas de origen y textos completos para comparar la salida."""
    wb = load_workbook(path, data_only=True)
    formulas = load_workbook(path, data_only=False)
    try:
        ws = wb["Hoja1"]
        if not text(ws["A1"].value).startswith("MEMORIAS 2025") or "PERSONAL" not in text(ws["A59"].value):
            raise ValueError("El archivo no corresponde a la estructura institucional de 2025")

        def value(row, col):
            return ws[f"{col}{row}"].value

        def records(start, end, columns):
            result = []
            for row in range(start, end + 1):
                data = {key: value(row, col) for key, col in columns.items()}
                substantive = [v for k, v in data.items() if k != "numero"]
                if not any(v not in (None, "", "-") for v in substantive):
                    continue
                data["source_row"] = row
                result.append(data)
            return result

        def paragraphs(start, end):
            return "\n\n".join(text(value(row, "A")) for row in range(start, end + 1) if text(value(row, "A")))

        data = {
            "grupo": {
                "nombre_unidad_academica": text(value(5, "A")).split(":", 1)[1].strip(),
                "nombre_sigla_grupo": text(value(6, "A")).split(":", 1)[1].strip(),
                "mail": text(value(9, "A")).split(":", 1)[1].strip(),
                "objetivo_desarrollo": paragraphs(47, 56),
            },
            "directivos": [{"cargo": cargo, "nombre_apellido": text(value(row, "A")).split(":", 1)[1].strip()}
                           for row, cargo in ((7, "Director"), (8, "Vicedirector"))],
            "organigrama": paragraphs(17, 44),
            "programa_actividades": paragraphs(322, 357),
            "investigadores": records(62, 67, {"nombre_apellido": "B", "categoria": "C", "incentivos": "D", "dedicacion": "E", "horas": "F"}),
            "personal": records(72, 73, {"nombre_apellido": "B", "horas": "C"}),
            "becarios": [],
            "equipamiento": records(126, 128, {"denominacion": "B", "fecha": "C", "monto": "D", "descripcion": "E"}),
            "proyectos": records(139, 141, {"tipo": "B", "codigo": "C", "periodo": "D", "monto": "E", "nombre": "F", "descripcion": "G", "logros": "H", "dificultades": "I", "fuente": "J"}),
            "participaciones": [{"source_row": row, "descripcion": text(value(row, "A"))}
                                for row in range(151, 208) if text(value(row, "A")) and row not in (191, 196, 205)],
            "reuniones": [],
            "revistas": records(227, 229, {"nombre_revista": "B", "pais": "C", "editorial": "D", "issn": "E", "titulo": "F"}),
            "articulos": [{"source_row": row, "descripcion": text(value(row, "A"))}
                          for row in range(240, 249) if text(value(row, "A"))],
            "docencia": records(267, 272, {"nombre_apellido": "B", "curso": "C", "posgrado": "E"}),
            "transferencias": [],
            "finanzas": records(309, 314, {"fuente": "B", "ingresos": "C", "egresos": "E"}),
            "formulas": {cell.coordinate: cell.value for row in formulas["Hoja1"] for cell in row if cell.data_type == "f"},
        }
        for start, end, category in ((81, 83, "Doctorado"), (86, 89, "Maestría/Especialización"), (93, 94, "Graduado"),
                                     (98, 104, "Alumno"), (108, 116, "Pasante"), (120, 122, "Tesina")):
            rows = records(start, end, {"nombre_apellido": "B", "fuente": "C", "horas": "D"})
            data["becarios"].extend({**item, "categoria": category} for item in rows if item["nombre_apellido"])
        for start, end, category in ((212, 217, "Nacional"), (221, 222, "Internacional")):
            data["reuniones"].extend({**item, "tipo": category} for item in records(start, end, {
                "nombre_reunion": "B", "procedencia": "C", "fecha": "D", "expositores": "E", "titulo": "F"}))
        for start, end, category in ((278, 281, "Transferencia de tecnología"), (284, 286, "I+D+i"),
                                     (289, 290, "Transferencia de conocimientos"), (293, 295, "Asistencia técnica o consultoría"),
                                     (298, 299, "Servicios técnicos"), (302, 304, "Difusión a la comunidad")):
            rows = records(start, end, {"denominacion": "B", "adoptante": "D", "demandante": "F", "monto": "G", "descripcion": "H"})
            data["transferencias"].extend({**item, "tipo": category} for item in rows if item["denominacion"] not in (None, "", "-"))
        for row in (309,):
            for col in ("C", "E"):
                if value(row, col) is None:
                    raise ValueError(f"La fórmula {col}{row} no conserva su resultado; recalcular el archivo en Excel")
        return data
    finally:
        wb.close()
        formulas.close()
