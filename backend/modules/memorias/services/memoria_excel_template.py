"""Completa la plantilla institucional sin reconstruir su maquetación."""
from copy import copy
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from math import ceil
from textwrap import wrap

from openpyxl.cell.cell import MergedCell
from openpyxl.styles.numbers import is_date_format
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange

EMPTY = "No registra datos en este período."
CONTENT_FONT = Font(name="Calibri", size=11, color="FF000000")


class TemplateWriter:
    def __init__(self, workbook):
        self.workbook = workbook
        self.sheet = workbook["Hoja1"]
        self.source = workbook.copy_worksheet(self.sheet)
        self.insertions = []

    def row(self, original):
        return original + sum(amount for position, amount in self.insertions if position <= original)

    def clear(self, start, end):
        for cells in self.sheet.iter_rows(min_row=self.row(start), max_row=self.row(end), max_col=12):
            for cell in cells:
                if not isinstance(cell, MergedCell):
                    cell.value = None
                    cell.comment = None

    def set(self, original_row, column, value, *, content=True):
        return self.set_at(self.row(original_row), column, value, content=content)

    def set_at(self, row, column, value, *, content=True):
        cell = self.sheet[f"{column}{row}"]
        if isinstance(value, (date, datetime)) and not is_date_format(cell.number_format):
            value = value.strftime("%d/%m/%Y")
        style = copy(cell._style)
        cell.value = value
        cell._style = style
        if isinstance(value, str):
            cell.data_type = "s"
        if content and value is not None:
            cell.font = copy(CONTENT_FONT)
            self.fit_text(cell)
        return cell

    def fit_text(self, cell):
        if not isinstance(cell.value, str) or cell.value == EMPTY or not cell.alignment.wrap_text:
            return
        merged = next((item for item in self.sheet.merged_cells.ranges if cell.coordinate in item), None)
        first, last = (merged.min_col, merged.max_col) if merged else (cell.column, cell.column)
        width = sum(self.sheet.column_dimensions[get_column_letter(column)].width or 13 for column in range(first, last + 1))
        lines = sum(max(1, len(wrap(line, width=max(10, int(width * 0.85))))) for line in cell.value.split("\n"))
        needed = ceil(lines * 15 + 5)
        final_row = merged.max_row if merged else cell.row
        heights = [self.sheet.row_dimensions[row].height or self.sheet.sheet_format.defaultRowHeight or 15
                   for row in range(cell.row, final_row + 1)]
        if needed > sum(heights):
            self.sheet.row_dimensions[cell.row].height = min(409, heights[0] + needed - sum(heights))

    def insert(self, before, amount, prototype):
        if not amount:
            return
        at = self.row(before)
        ranges = []
        for merged in self.sheet.merged_cells.ranges:
            updated = CellRange(str(merged))
            if updated.min_row >= at:
                updated.shift(row_shift=amount)
            elif updated.max_row >= at:
                updated.max_row += amount
            ranges.append(str(updated))
        dimensions = {index: copy(dimension) for index, dimension in self.sheet.row_dimensions.items()}
        self.sheet.insert_rows(at, amount)
        self.sheet.row_dimensions.clear()
        for index, dimension in dimensions.items():
            new_index = index + amount if index >= at else index
            dimension.index = new_index
            self.sheet.row_dimensions[new_index] = dimension
        self.sheet.merged_cells = MultiCellRange()
        for merged in ranges:
            self.sheet.merge_cells(merged)
        for target in range(at, at + amount):
            self.clone_row(prototype, target)
        for page_break in self.sheet.row_breaks.brk:
            if page_break.id >= at:
                page_break.id += amount
        self.insertions.append((before, amount))

    def clone_row(self, prototype, target):
        dimension = copy(self.source.row_dimensions[prototype])
        dimension.index = target
        self.sheet.row_dimensions[target] = dimension
        for column in range(1, 13):
            self.sheet.cell(target, column)._style = copy(self.source.cell(prototype, column)._style)
        for merged in self.source.merged_cells.ranges:
            if merged.min_row == prototype and merged.max_row == prototype:
                self.sheet.merge_cells(start_row=target, end_row=target,
                    start_column=merged.min_col, end_column=merged.max_col)

    def record_layout(self, row, columns, prototype, replace_placeholder=False):
        # Sustituir los mensajes combinados únicamente si deben entrar registros.
        if not replace_placeholder and not any(isinstance(self.sheet[f"{column}{row}"], MergedCell) for column in columns):
            return
        for merged in list(self.sheet.merged_cells.ranges):
            if merged.min_row <= row <= merged.max_row:
                self.sheet.unmerge_cells(str(merged))
        self.clone_row(prototype, row)

    def table(self, start, end, columns, records, *, prototype=None, empty_row=None, replace_placeholder=False):
        records = list(records)
        prototype = prototype or start
        self.clear(start, end)
        self.insert(end + 1, max(0, len(records) - (end - start + 1)), prototype)
        if not records:
            self.set(empty_row or start, columns[0], EMPTY)
            return
        first = self.row(start)
        for index, values in enumerate(records):
            row = first + index
            self.record_layout(row, columns, prototype, replace_placeholder)
            for column, value in zip(columns, values):
                self.set_at(row, column, value)

    def narrative(self, start, end, paragraphs, *, empty_row=None):
        paragraphs = [str(value) for value in paragraphs if value not in (None, "")]
        anchors = [row for row in range(start, end + 1)
                   if not isinstance(self.source[f"A{row}"], MergedCell)
                   and isinstance(self.source[f"A{row}"].value, str)
                   and self.source[f"A{row}"].value.strip()]
        if not anchors:
            anchors = [empty_row or start]
        self.clear(start, end)
        if not paragraphs:
            self.set(empty_row or anchors[0], "A", EMPTY)
            return
        if len(anchors) == 1:
            if len(paragraphs) > 1:
                original = self.source[f"A{anchors[0]}"]
                merged = next((item for item in self.source.merged_cells.ranges if original.coordinate in item), None)
                prototype = 144 if merged and merged.max_row > merged.min_row else anchors[0]
                self.table(anchors[0], end, ["A"], [[paragraph] for paragraph in paragraphs],
                           prototype=prototype, replace_placeholder=True)
                return
            self.set(anchors[0], "A", "\n\n".join(paragraphs))
            return
        self.insert(end + 1, max(0, len(paragraphs) - len(anchors)), anchors[0])
        for index, paragraph in enumerate(paragraphs):
            target = self.row(anchors[index]) if index < len(anchors) else self.row(end) + 1 + index - len(anchors)
            self.set_at(target, "A", paragraph)

    def finish(self):
        self.workbook.remove(self.source)
        output = BytesIO()
        self.workbook.save(output)
        output.seek(0)
        return output


def money(item):
    value = item.get("monto_equivalente_ars")
    return Decimal(str(value if value is not None else item.get("monto") or 0))


def civil_date(value):
    if isinstance(value, str):
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return value
    return value


def render_memoria(workbook, sources, context, helpers):
    w = TemplateWriter(workbook)
    group = context["grupo"]
    memoria = sources["memoria"]
    year = memoria.periodo_fin.year
    sigla = group["nombre_sigla_grupo"].rsplit(" - ", 1)[-1]
    label = str(year) if memoria.periodo_inicio == date(year, 1, 1) and memoria.periodo_fin == date(year, 12, 31) else f"{helpers._format_date(memoria.periodo_inicio)} al {helpers._format_date(memoria.periodo_fin)}"
    w.set(1, "A", f"MEMORIAS {label} DEL GRUPO UTN - {sigla}", content=False)
    directivos = context.get("directivos", [])

    def authority(cargo):
        matching = [item for item in directivos if (item.get("cargo") or "").strip().casefold() == cargo.casefold()]
        matching.sort(key=lambda item: item.get("fecha_inicio") or "")
        return matching[-1]["nombre_apellido"] if matching else "-"

    for row, value in ((5, group["nombre_unidad_academica"]), (6, group["nombre_sigla_grupo"]), (7, authority("Director")),
                       (8, authority("Vicedirector")), (9, group["mail"])):
        prefix = w.source[f"A{row}"].value.split(":", 1)[0]
        w.set(row, "A", f"{prefix}: {value}", content=False)
    consejo = [item for item in directivos if "consejo" in (item.get("cargo") or "").casefold()]
    w.table(12, 15, ["A", "B", "D"], [[i, item["nombre_apellido"], item.get("cargo") or "-"] for i, item in enumerate(consejo, 1)])
    # El organigrama pertenece al formato institucional; un informe de
    # actividades UCT no se reinterpreta como estructura organizacional.
    w.narrative(47, 56, [group.get("objetivo_desarrollo")])
    w.table(62, 67, ["A", "B", "C", "D", "E", "F"], [[i, item.get("nombre_apellido"), item.get("categoria_utn_nombre") or "-",
        item.get("programa_incentivos_nombre") or "-", item.get("tipo_dedicacion_nombre") or "-", item.get("horas_semanales")]
        for i, item in enumerate(sources["investigadores"], 1)])
    for start, end, category in ((72, 73, "profesional"), (77, 77, "apoyo")):
        items = [item for item in sources["personal"] if helpers._clasificar_personal(item.get("tipo_personal_nombre")) == category]
        w.table(start, end, ["A", "B", "C"], [[i, item.get("nombre_apellido"), item.get("horas_semanales")] for i, item in enumerate(items, 1)])
    for start, end, category in ((81, 83, "doctorado"), (86, 89, "maestria"), (93, 94, "graduado"),
                                  (98, 104, "alumno"), (108, 116, "pasante"), (120, 122, "tesina")):
        items = [item for item in sources["becarios"] if helpers._clasificar_becario(item.get("tipo_formacion_nombre")) == category]
        w.table(start, end, ["A", "B", "C", "D"], [[i, item.get("nombre_apellido"), item.get("fuentes_financiamiento_beca") or "-", item.get("horas_semanales")]
            for i, item in enumerate(items, 1)], prototype=86 if category == "doctorado" else start,
            empty_row=82 if category == "doctorado" else None, replace_placeholder=category == "doctorado")
    w.table(126, 128, ["A", "B", "C", "D", "E"], [[i, item.get("denominacion"), civil_date(item.get("fecha_incorporacion")),
        Decimal(str(item.get("monto_invertido") or 0)), item.get("descripcion_breve")] for i, item in enumerate(sources["equipamiento"], 1)])
    w.table(132, 133, ["A", "B", "C", "D", "E", "F"], [[i, item.get("titulo"), helpers._join_dict_names(item.get("autores"), "nombre_apellido"),
        item.get("editorial"), item.get("anio"), civil_date(item.get("fecha"))] for i, item in enumerate(sources["documentacion"], 1)])
    w.table(139, 141, ["A", "B", "C", "D", "F", "G", "H", "I", "J"], [[i, item.get("tipo_proyecto_nombre"), item.get("codigo_proyecto"),
        helpers._format_period(civil_date(item.get("fecha_inicio")), civil_date(item.get("fecha_fin"))), item.get("nombre_proyecto"),
        item.get("descripcion_proyecto"), item.get("logros_obtenidos") or context.get("logros_proyectos", {}).get(str(item.get("proyecto_investigacion_id")), EMPTY),
        item.get("dificultades_proyecto") or "-", item.get("fuente_financiamiento_nombre") or "-"] for i, item in enumerate(sources["proyectos"], 1)])
    w.narrative(144, 145, [f"{item.get('descripcion')} ({helpers._format_date(civil_date(item.get('fecha')))})" for item in sources["distinciones"]])
    w.narrative(147, 147, [])
    sections = {"otras": [], "cursos": [], "gestion": [], "eventos": []}
    for item in sources["participaciones"]:
        value = f"{item.get('nombre_evento') or '-'}\n{item.get('forma_participacion') or '-'}"
        kind = value.casefold()
        category = "gestion" if "gestión interna" in kind or "gestion interna" in kind or "asamblea" in kind else (
            "eventos" if "organización" in kind or "organizacion" in kind else "cursos" if "curso" in kind or "taller" in kind else "otras")
        sections[category].append(value)
    w.narrative(150, 190, sections["otras"])
    w.narrative(192, 195, sections["cursos"])
    w.narrative(197, 204, sections["gestion"])
    w.narrative(206, 208, sections["eventos"])
    visitas = [f"{item.get('procedencia') or '-'}: {item.get('razon') or '-'}" for item in sources["visitas"]]
    if visitas:
        w.set(148, "A", "6.2.- Visitantes del país y del extranjero:", content=False)
        w.insert(149, len(visitas), 144)
        first = w.row(149) - len(visitas)
        for index, paragraph in enumerate(visitas):
            w.set_at(first + index, "A", paragraph)
    for start, end, category in ((212, 217, "nacional"), (221, 222, "internacional")):
        items = [item for item in sources["trabajos_reunion"] if helpers._clasificar_trabajo_reunion(item.get("tipo_reunion_nombre")) == category]
        w.table(start, end, ["A", "B", "C", "D", "E", "F"], [[i, item.get("nombre_reunion"), item.get("procedencia"),
            civil_date(item.get("fecha_presentacion")), helpers._autores_texto(item.get("autores", [])),
            (item.get("titulo_trabajo") or "-") + ("\nEnlace: " + item["enlace"] if item.get("enlace") else "")]
            for i, item in enumerate(items, 1)])
    revistas = sources["trabajos_revista"]
    w.table(227, 229, ["A", "B", "C", "D", "E", "F"], [[i, item.get("nombre_revista"), item.get("pais"), item.get("editorial"), item.get("issn"),
        (item.get("titulo_trabajo") or "-") + "\nAutores: " + helpers._autores_texto(item.get("autores", [])) +
        ("\nEnlace: " + item["enlace"] if item.get("enlace") else "")] for i, item in enumerate(revistas, 1)])
    w.narrative(233, 237, [f"{item.get('titulo_trabajo')}\n{item.get('nombre_revista')}\n{item.get('issn')}" for item in revistas
                         if "libro" in (item.get("tipo_revista_nombre") or "").casefold() or "ISBN" in (item.get("issn") or "").upper()])
    w.narrative(239, 249, [item.get("descripcion") or item.get("titulo") for item in sources["articulos"]], empty_row=240)
    w.narrative(251, 255, [item.get("nombre_articulo") for item in sources["registros"] if "patente" in (item.get("tipo_registro_nombre") or "").casefold()])
    for start, end, category in ((258, 260, "intelectual"), (262, 264, "industrial")):
        items = [item for item in sources["registros"] if "patente" not in (item.get("tipo_registro_nombre") or "").casefold()
                 and helpers._clasificar_registro(item.get("tipo_registro_nombre")) == category]
        w.narrative(start, end, [f"{item.get('nombre_articulo')} - {item.get('organismo_registrante')} - {helpers._format_date(civil_date(item.get('fecha_registro')))}" for item in items])
    w.table(267, 272, ["A", "B", "C", "E"], [[i, item.get("investigador_nombre"), item.get("grado_academico_nombre") or "-",
        f"{item.get('curso') or '-'}\n{item.get('institucion') or '-'}"] for i, item in enumerate(sources["actividades"], 1)])
    for start, end, category in ((278, 279, "tecnologia"), (284, 285, "idi"), (289, 289, "conocimientos"),
                                (293, 294, "asistencia"), (298, 298, "servicios"), (302, 304, "difusion")):
        items = [item for item in sources["transferencias"] if helpers._clasificar_transferencia(item.get("tipo_contrato_nombre")) == category]
        w.table(start, end, ["A", "B", "D", "F", "G", "H"], [[i, item.get("denominacion"), helpers._join_dict_names(item.get("adoptantes"), "adoptante_nombre"),
            item.get("demandante"), Decimal(str(item.get("monto") or 0)), item.get("descripcion_actividad")] for i, item in enumerate(items, 1)])
    for start, end, category in ((309, 314, "CORRIENTE"), (318, 319, "CAPITAL")):
        totals = {}
        for item in sources["erogaciones"]:
            ingreso = item.get("tipo_movimiento") == "INGRESO"
            if (ingreso and category != "CORRIENTE") or (not ingreso and item.get("categoria_erogacion_codigo") != category):
                continue
            source = item.get("fuente_financiamiento_nombre") or "-"
            entry = totals.setdefault(source, [Decimal(0), Decimal(0)])
            entry[0 if ingreso else 1] += money(item)
        w.table(start, end, ["A", "B", "C", "E"], [[i, name, values[0], values[1]] for i, (name, values) in enumerate(totals.items(), 1)])
    w.set(321, "A", f"VI - PROGRAMA DE ACTIVIDADES para {year + 1}", content=False)
    w.narrative(322, 370, (context.get("programa_actividades") or "").split("\n\n"))
    return w.finish()
